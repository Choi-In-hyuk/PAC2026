"""Generate a Word document summarizing the mesh->CSV preprocessing guide."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_h(doc, text, level=1):
    doc.add_heading(text, level=level)


def add_p(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    return p


def add_code(doc, code):
    p = doc.add_paragraph()
    r = p.add_run(code)
    r.font.name = "Courier New"
    r.font.size = Pt(9)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()


doc = Document()

# Title
title = doc.add_heading("메쉬 → surface CSV 변환 가이드", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_p(doc, "arm_coverage (ICRA'25) 알고리즘 입력 데이터 전처리", italic=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph()

# 전체 흐름
add_h(doc, "1. 전체 흐름", 1)
add_code(doc,
"입력                                    출력\n"
"─────                                   ─────\n"
"car.stl        ────┐\n"
"tool size (cm) ────┤  ──► [변환 스크립트] ──►  surfaces/<robot>_<ts>_<task>.csv\n"
"work region    ────┤\n"
"robot base pose ───┘\n")

# Step 0
add_h(doc, "2. Step 0: 사전 결정 사항", 1)
add_p(doc, "코드를 짜기 전에 반드시 결정해야 할 설정값들입니다.")
add_table(doc,
          ["항목", "예시 값", "왜 필요한가"],
          [
              ("로봇", "Franka Panda", "URDF, base/EE link 이름"),
              ("차체 메쉬 파일", "car_hood.stl", "입력"),
              ("작업 영역 (ROI)", "bbox x∈[1,2], y∈[-0.8,0.8]", "메쉬 중 어디 쓸지"),
              ("공구 크기", "샌딩 디스크 지름 10cm", "샘플 간격 결정"),
              ("샘플 간격", "공구 크기 × 0.7 ≈ 7cm", "overlap 30%"),
              ("로봇 base의 월드 위치", "(0, 0, 0) + 회전 [0,0,0]", "좌표 변환"),
              ("차체의 월드 위치", "(1.5, 0, 0.5)", "좌표 변환"),
              ("작업 종류", "샌딩", "tolerance 결정"),
              ("tolerance", "[0.03, 0.03, 0, 0, 0, 999]", "알고리즘에 전달"),
          ])

# Step 1
add_h(doc, "3. Step 1: 메쉬 로드", 1)
add_p(doc, "입력 정보: 차체 STL/OBJ/PLY 파일 (CAD 또는 3D 스캔)")
add_p(doc, "필요한 라이브러리: trimesh")
add_p(doc, "얻어지는 데이터:")
for b in ["vertices: (V, 3) — 메쉬 정점들",
          "faces: (F, 3) — 삼각형 (vertex index 3개)",
          "face_normals: (F, 3) — 각 삼각형의 법선 (자동 계산)",
          "face_centers: (F, 3) — 각 삼각형 중심"]:
    doc.add_paragraph(b, style="List Bullet")
add_code(doc,
"import trimesh\n"
"mesh = trimesh.load(\"car_hood.stl\")\n"
"# mesh.vertices, mesh.faces, mesh.face_normals 다 자동으로 생김\n")
add_p(doc, "체크해야 할 것:", bold=True)
for b in ["메쉬 단위가 뭔지 (mm? m?) → ROS 관례는 m, 아니면 mesh.apply_scale(0.001)",
          "법선 방향이 바깥쪽을 가리키는지 → mesh.fix_normals()로 교정",
          "메쉬가 watertight인지 (닫혀 있는지) → mesh.is_watertight 확인"]:
    doc.add_paragraph(b, style="List Bullet")

# Step 2
add_h(doc, "4. Step 2: 작업 영역 크롭 (ROI 필터링)", 1)
add_p(doc, "입력 정보: 작업 영역 bbox 또는 평면/실린더 등, 또는 로봇 reach radius")
add_p(doc, "방법 A: Bounding box 안의 삼각형만", bold=True)
add_code(doc,
"bbox_min = [1.0, -0.5, 0.3]\n"
"bbox_max = [2.0,  0.5, 0.7]\n"
"mask = np.all((mesh.triangles_center >= bbox_min) & \n"
"              (mesh.triangles_center <= bbox_max), axis=1)\n"
"roi_mesh = mesh.submesh([np.where(mask)[0]], append=True)\n")
add_p(doc, "방법 B: 평면으로 자르기", bold=True)
add_code(doc,
"roi_mesh = mesh.slice_plane(plane_origin=[0,0,0.5], plane_normal=[0,0,1])\n")
add_p(doc, "방법 C: 법선 방향 필터 (예: 위를 바라보는 face만)", bold=True)
add_code(doc,
"mask = mesh.face_normals @ [0,0,1] > 0.5  # 위 반구만\n")
add_p(doc, "실제로는 A+C 조합이 자주 쓰임 (영역 + 법선 방향).")

# Step 3
add_h(doc, "5. Step 3: 표면 샘플링 (알고리즘 선택)", 1)
add_p(doc, "입력 정보: 크롭된 메쉬 / 샘플 개수 N (또는 샘플 간격) / 샘플링 알고리즘 선택")
add_p(doc, "얻어지는 데이터:")
for b in ["points: (N, 3) — 표면 위 점",
          "face_idx: (N,) — 각 점이 어느 삼각형에서 나왔는지",
          "normals: (N, 3) = mesh.face_normals[face_idx]"]:
    doc.add_paragraph(b, style="List Bullet")
add_h(doc, "샘플링 알고리즘 비교", 2)
add_table(doc,
          ["방법", "함수", "특징", "추천?"],
          [
              ("Uniform random", "trimesh.sample.sample_surface", "랜덤, 뭉침 발생", "✗"),
              ("Poisson-disk", "trimesh.sample.sample_surface_even", "균등 간격 보장", "★ 기본"),
              ("Zigzag/Isoline", "직접 작성", "쓸고 지나가는 패턴", "샌딩/와이핑에 더 좋음"),
              ("Grid projection", "2D 격자 투영", "평면 작업용", "평면이면 추천"),
          ])
add_p(doc, "기본 추천: Poisson-disk", bold=True)
add_code(doc,
"tool_diameter = 0.10  # 10cm\n"
"sample_radius = tool_diameter * 0.7  # 7cm\n\n"
"points, face_idx = trimesh.sample.sample_surface_even(roi_mesh,\n"
"                                                     count=1000,\n"
"                                                     radius=sample_radius)\n"
"normals = roi_mesh.face_normals[face_idx]\n")
add_p(doc, "트레이드오프: 처음엔 Poisson-disk로 돌려보고, 결과가 지저분하면 zigzag로 바꾸는 걸 추천.")

# Step 4
add_h(doc, "6. Step 4: 좌표계 변환 (메쉬 → 로봇 base)", 1)
add_p(doc, "입력 정보: points/normals (메쉬 좌표계) + T_robot_from_mesh 4×4 변환")
add_code(doc,
"import numpy as np\n"
"from scipy.spatial.transform import Rotation as R\n\n"
"# 월드에서 차체 위치/자세\n"
"T_world_from_mesh = np.eye(4)\n"
"T_world_from_mesh[:3, 3] = [1.5, 0.0, 0.5]\n"
"T_world_from_mesh[:3, :3] = R.from_euler('z', 0).as_matrix()\n\n"
"# 월드에서 로봇 base 위치/자세\n"
"T_world_from_robot = np.eye(4)\n"
"T_world_from_robot[:3, 3] = [0, 0, 0]\n\n"
"# 로봇 base에서 본 차체\n"
"T_robot_from_mesh = np.linalg.inv(T_world_from_robot) @ T_world_from_mesh\n\n"
"# 적용\n"
"points_h = np.hstack([points, np.ones((len(points),1))])\n"
"points_robot = (T_robot_from_mesh @ points_h.T).T[:, :3]\n"
"normals_robot = (T_robot_from_mesh[:3,:3] @ normals.T).T\n")
add_p(doc, "실제 셋업에선 캘리브레이션 필요 (ArUco 마커, hand-eye calibration 등). 시뮬레이션/프로토타입이면 숫자로 결정.")

# Step 5
add_h(doc, "7. Step 5: 법선 → 쿼터니언 변환", 1)
add_p(doc, "입력 정보: normals_robot (로봇 좌표계) + 공구 축 convention (EE Z축이 공구 방향)")
add_p(doc, "목표: EE Z축을 −법선 방향(공구가 표면을 눌러 들어가는 방향)으로 세우기")
add_code(doc,
"from scipy.spatial.transform import Rotation as R\n\n"
"def normals_to_quats(normals, roll_angle=0.0):\n"
"    quats = np.zeros((len(normals), 4))\n"
"    z_default = np.array([0, 0, 1])\n"
"    for i, n in enumerate(normals):\n"
"        z_target = -n / np.linalg.norm(n)\n"
"        rot, _ = R.align_vectors([z_target], [z_default])\n"
"        if roll_angle != 0:\n"
"            roll = R.from_rotvec(z_target * roll_angle)\n"
"            rot = roll * rot\n"
"        quats[i] = rot.as_quat()  # xyzw\n"
"    return quats\n")
add_p(doc, "체크 포인트:", bold=True)
for b in ["scipy의 as_quat()는 xyzw 순서 (레포 CSV도 xyzw)",
          "법선과 EE Z축이 평행이면 align_vectors가 degenerate — scipy가 handling",
          "완전 반대 방향이면 180° 회전, 축은 임의 — scipy가 알아서 처리"]:
    doc.add_paragraph(b, style="List Bullet")

# Step 6
add_h(doc, "8. Step 6: CSV 저장", 1)
add_code(doc,
"import csv\n"
"from datetime import datetime\n\n"
"def save_surface_csv(points, quats, robot_name, task_name, out_dir=\"surfaces\"):\n"
"    ts = datetime.now().strftime(\"%Y%m%d_%H%M%S\")\n"
"    fname = f\"{out_dir}/{robot_name}_{ts}_{task_name}.csv\"\n"
"    with open(fname, \"w\", newline=\"\") as f:\n"
"        w = csv.writer(f)\n"
"        w.writerow([\"time\", \"target-POS_X\", \"target-POS_Y\", \"target-POS_Z\",\n"
"                             \"target-ROT_X\", \"target-ROT_Y\", \"target-ROT_Z\", \"target-ROT_W\"])\n"
"        for i in range(len(points)):\n"
"            w.writerow([i, *points[i], *quats[i]])\n")
add_p(doc, "파일명 규칙:", bold=True)
for b in ["<robot> → panda / ur5 / ur5sander / 본인 로봇 이름",
          "<task> → floor / wok / wokreal / 새로 추가할 작업명 (예: hood, doorpanel)"]:
    doc.add_paragraph(b, style="List Bullet")

# Step 7
add_h(doc, "9. Step 7: run_experiments.py에 분기 추가", 1)
add_p(doc, "CSV를 만들어도 알고리즘이 tolerance를 모름. 코드 한 줄 추가.")
add_p(doc, "편집할 파일: scripts/run_experiments.py", bold=True)
add_code(doc,
"if task == \"wok\":\n"
"    tolerances = [0.05, 0.05, 0., 0., 0., 999.]\n"
"elif task == \"wokreal\":\n"
"    tolerances = [0., 0., 0., 0., 0., 999.]\n"
"elif task == \"floor\":\n"
"    tolerances = [0., 0., 0., 0., 0., 0.]\n"
"elif task == \"hood\":                          # ← 추가\n"
"    tolerances = [0.03, 0.03, 0., 0., 0., 999.]  # 샌딩용\n")

# 전체 스크립트 뼈대
add_h(doc, "10. 전체 스크립트 뼈대 (약 100줄)", 1)
add_code(doc,
"\"\"\"\n"
"mesh_to_surface.py\n"
"Usage: python mesh_to_surface.py car.stl --robot ur5 --task hood\n"
"\"\"\"\n"
"import numpy as np, csv, argparse\n"
"from datetime import datetime\n"
"from scipy.spatial.transform import Rotation as R\n"
"import trimesh\n\n"
"def crop_mesh(mesh, bbox_min, bbox_max):\n"
"    mask = np.all((mesh.triangles_center >= bbox_min) & \n"
"                  (mesh.triangles_center <= bbox_max), axis=1)\n"
"    return mesh.submesh([np.where(mask)[0]], append=True)\n\n"
"def sample_points(mesh, radius):\n"
"    pts, face_idx = trimesh.sample.sample_surface_even(mesh,\n"
"                                                       count=10000,\n"
"                                                       radius=radius)\n"
"    normals = mesh.face_normals[face_idx]\n"
"    return pts, normals\n\n"
"def build_transform(pos, rpy):\n"
"    T = np.eye(4)\n"
"    T[:3, 3] = pos\n"
"    T[:3, :3] = R.from_euler('xyz', rpy).as_matrix()\n"
"    return T\n\n"
"def transform_points(T, pts):\n"
"    h = np.hstack([pts, np.ones((len(pts),1))])\n"
"    return (T @ h.T).T[:, :3]\n\n"
"def transform_vectors(T, vecs):\n"
"    return (T[:3,:3] @ vecs.T).T\n\n"
"def normals_to_quats(normals):\n"
"    z_default = np.array([0, 0, 1])\n"
"    quats = np.zeros((len(normals), 4))\n"
"    for i, n in enumerate(normals):\n"
"        z_target = -n / np.linalg.norm(n)\n"
"        rot, _ = R.align_vectors([z_target], [z_default])\n"
"        quats[i] = rot.as_quat()\n"
"    return quats\n\n"
"def save_csv(points, quats, path):\n"
"    with open(path, \"w\", newline=\"\") as f:\n"
"        w = csv.writer(f)\n"
"        w.writerow([\"time\", \"target-POS_X\",\"target-POS_Y\",\"target-POS_Z\",\n"
"                             \"target-ROT_X\",\"target-ROT_Y\",\"target-ROT_Z\",\"target-ROT_W\"])\n"
"        for i in range(len(points)):\n"
"            w.writerow([i, *points[i], *quats[i]])\n\n"
"def main():\n"
"    ap = argparse.ArgumentParser()\n"
"    ap.add_argument(\"mesh\")\n"
"    ap.add_argument(\"--robot\", required=True)\n"
"    ap.add_argument(\"--task\", required=True)\n"
"    ap.add_argument(\"--sample-radius\", type=float, default=0.07)\n"
"    ap.add_argument(\"--mesh-pos\", nargs=3, type=float, default=[1.5, 0, 0.5])\n"
"    ap.add_argument(\"--mesh-rpy\", nargs=3, type=float, default=[0, 0, 0])\n"
"    ap.add_argument(\"--bbox-min\", nargs=3, type=float, default=None)\n"
"    ap.add_argument(\"--bbox-max\", nargs=3, type=float, default=None)\n"
"    args = ap.parse_args()\n\n"
"    mesh = trimesh.load(args.mesh)\n"
"    mesh.fix_normals()\n\n"
"    if args.bbox_min and args.bbox_max:\n"
"        mesh = crop_mesh(mesh, args.bbox_min, args.bbox_max)\n\n"
"    pts, normals = sample_points(mesh, args.sample_radius)\n\n"
"    T = build_transform(args.mesh_pos, args.mesh_rpy)\n"
"    pts_r = transform_points(T, pts)\n"
"    normals_r = transform_vectors(T, normals)\n\n"
"    quats = normals_to_quats(normals_r)\n\n"
"    ts = datetime.now().strftime(\"%Y%m%d_%H%M%S\")\n"
"    out = f\"surfaces/{args.robot}_{ts}_{args.task}.csv\"\n"
"    save_csv(pts_r, quats, out)\n"
"    print(f\"Saved {len(pts_r)} points to {out}\")\n\n"
"if __name__ == \"__main__\":\n"
"    main()\n")

# 요약 표
add_h(doc, "11. 요약: 각 단계에서 뽑는 정보", 1)
add_table(doc,
          ["단계", "입력 정보", "뽑아내는 정보", "핵심 함수"],
          [
              ("1. 메쉬 로드", "STL/OBJ 파일", "vertices, faces, face_normals", "trimesh.load()"),
              ("2. ROI 크롭", "bbox 또는 평면", "작업 영역 메쉬", "mask + submesh()"),
              ("3. 샘플링", "공구 크기", "points, face_idx", "sample_surface_even()"),
              ("3b. 법선", "face_idx", "normals", "mesh.face_normals[face_idx]"),
              ("4. 좌표 변환", "로봇/차체 상대 위치", "points_robot, normals_robot", "행렬 곱"),
              ("5. 쿼터니언", "normals_robot", "quats (xyzw)", "R.align_vectors()"),
              ("6. CSV 저장", "points, quats", "파일", "csv.writer"),
              ("7. tolerance", "작업 성격", "분기", "코드 한 줄"),
          ])

# 다음 액션
add_h(doc, "12. 다음 액션 체크리스트", 1)
for i, item in enumerate([
    "차체 메쉬 입수 (팀에서 받기, 또는 CAD에서 export)",
    "로봇 결정 + base/차체 상대 위치 정의",
    "tolerance 결정 (작업 성격)",
    "위 스크립트 짜서 CSV 생성",
    "run_experiments.py tolerance 분기 추가",
    "알고리즘 돌리기 → motions/*.csv",
    "validate + visualize",
], 1):
    doc.add_paragraph(f"{i}. {item}", style="List Number")

out_path = "/home/choi/pac/arm_coverage/docs/mesh_to_csv_guide.docx"
import os
os.makedirs(os.path.dirname(out_path), exist_ok=True)
doc.save(out_path)
print(f"saved {out_path}")
